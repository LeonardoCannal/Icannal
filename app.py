from flask import Flask, render_template, request, jsonify, send_file, send_from_directory, redirect, url_for, has_request_context
from flask.sessions import SecureCookieSessionInterface
from werkzeug.middleware.proxy_fix import ProxyFix
from urllib.parse import urlparse
from scraper import (
    buscar_medicos_doctoralia,
    buscar_medicos_sechat,
    buscar_medicos_apepi,
    CacheTTL,
    normalizar,
    buscar_medicos_fonte_generica,
    FONTES_GENERICAS,
)
from auth import (
    inicializar_auth,
    criar_usuario,
    buscar_usuario_por_email,
    buscar_usuario_por_cpf,
    buscar_usuario_por_id,
    cpf_valido,
    formatar_cpf,
    checar_senha,
    atualizar_ultimo_login,
    listar_usuarios_com_contagem_painel,
    usuario_logado,
    fazer_login,
    fazer_logout,
    login_required,
    admin_required,
    ErroIntegridade,
    computar_chave_medico,
    adicionar_ao_painel,
    mover_no_painel,
    listar_painel_usuario,
    chaves_no_painel_usuario,
    chaves_no_painel_todos,
    registrar_solicitacao_senha,
    listar_solicitacoes_senha_pendentes,
    contar_solicitacoes_senha_pendentes,
    marcar_solicitacao_atendida,
    buscar_solicitacao_senha_por_id,
    definir_senha,
    registrar_busca,
    estatisticas_dashboard,
    estatisticas_usuario,
    bucket_painel,
    listar_horarios_dia,
    agendar_medico,
    cancelar_agendamento,
    listar_compromissos_mes,
    salvar_observacoes,
    dias_sem_visitar,
    roteiro_do_dia,
    contar_visitas_amanha,
    RESULTADOS_VISITA,
    buscar_compromisso_do_medico,
    importar_base_propria,
    buscar_base_propria,
    contar_base_propria,
    admin_importar_painel_usuario,
    tornar_admin,
    eh_admin_mestre,
    excluir_usuario,
    remover_admin,
    remover_medico_painel,
    listar_cadastros_pendentes,
    contar_cadastros_pendentes,
    aprovar_usuario,
    limpar_medicos_inativos,
    limpar_medicos_inativos_todos_usuarios,
    listar_notificacoes_usuario,
    contar_notificacoes_nao_lidas,
    marcar_notificacoes_lidas,
    listar_medicos_removidos_usuario,
    listar_medicos_removidos_geral,
    data_valida,
)
import requests
import webbrowser
import threading
import gzip
from concurrent.futures import ThreadPoolExecutor, as_completed
import sys
import os
import io
import csv
import secrets
import time
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def resource_path(relative_path):
    """Acha o caminho certo dos arquivos SÓ DE LEITURA (templates, static),
    tanto rodando normal quanto dentro do .exe empacotado."""
    base_path = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base_path, relative_path)


def caminho_dados_persistentes(nome_arquivo):
    """
    Acha o caminho certo pra arquivos que precisam ser GRAVADOS e
    PERSISTIR entre execuções (como o banco de dados de usuários).
    Dentro do .exe, a pasta que o resource_path() usa é temporária e é
    apagada quando o programa fecha — então o banco de dados precisa
    ficar do lado do .exe, não dentro dela.
    """
    if getattr(sys, "frozen", False):
        base_path = os.path.dirname(sys.executable)
    else:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, nome_arquivo)


app = Flask(
    __name__,
    template_folder=resource_path("templates"),
    static_folder=resource_path("static"),
)

# A SECRET_KEY protege as sessões de login. Em produção (site no ar),
# defina a variável de ambiente SECRET_KEY com um valor fixo e secreto —
# senão, toda vez que o servidor reiniciar, todo mundo é deslogado.
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", secrets.token_hex(32))

# Cookie de login: não acessível por JavaScript e não enviado em requisições
# vindas de outros sites. Com o site em HTTPS (VPS/Render), defina a variável
# COOKIE_SEGURO=1 pra ele também só trafegar criptografado.
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.environ.get("COOKIE_SEGURO") == "1"


class _SessaoSegura(SecureCookieSessionInterface):
    """Liga o "cookie só por HTTPS" sozinho quando o acesso é por HTTPS
    (produção), sem quebrar testes locais por HTTP."""
    def get_cookie_secure(self, app):
        if app.config.get("SESSION_COOKIE_SECURE"):
            return True
        return has_request_context() and request.is_secure


app.session_interface = _SessaoSegura()

# Atrás do nginx (VPS) ou do proxy do Render, o Flask enxerga todo mundo como
# vindo de 127.0.0.1 e por HTTP. Isso faz ele ler o IP real do visitante e
# saber que o acesso foi por HTTPS (usado nos limites por IP e no cookie).
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

# Nenhum envio maior que isso (planilhas de até 10 MB passam folgado).
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024


@app.before_request
def _bloquear_envio_de_outros_sites():
    """Defesa extra contra um site malicioso tentar enviar ações (salvar,
    apagar, aprovar...) em nome de quem está logado: ações só são aceitas
    quando vêm das páginas do próprio i.cannal."""
    if request.method in ("GET", "HEAD", "OPTIONS"):
        return None
    origem = request.headers.get("Origin") or request.headers.get("Referer")
    if not origem or origem == "null":
        return None
    host_origem = urlparse(origem).netloc.lower()
    if host_origem and host_origem != request.host.lower():
        return jsonify({"ok": False, "erro": "Requisição recusada."}), 403
    return None


@app.after_request
def _cabecalhos_de_seguranca(resposta):
    h = resposta.headers
    h.setdefault("X-Content-Type-Options", "nosniff")
    h.setdefault("X-Frame-Options", "DENY")
    h.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    h.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=(), payment=()")
    h.setdefault("Content-Security-Policy", "frame-ancestors 'none'; base-uri 'self'; object-src 'none'; form-action 'self'")
    if request.is_secure:
        h.setdefault("Strict-Transport-Security", "max-age=31536000")
    # Páginas e dados com informação de usuário: nenhum cache compartilhado
    # (proxy, CDN) pode guardar; o navegador sempre confere se mudou.
    if not request.path.startswith("/static/") and "Cache-Control" not in h:
        h["Cache-Control"] = "private, no-cache"
    return resposta


# ---------------------------------------------------------------------------
# Páginas mais leves
# ---------------------------------------------------------------------------

# CSS, ícones e manifest ganham um "?v=..." automático que muda quando o
# arquivo muda. Assim o navegador pode guardar esses arquivos por 30 dias
# sem medo: depois de uma atualização, o endereço muda e ele baixa de novo.
_VERSOES_ESTATICOS = {}


@app.url_defaults
def _versao_dos_estaticos(endpoint, valores):
    if endpoint != "static" or "filename" not in valores or "v" in valores:
        return
    nome = valores["filename"]
    versao = _VERSOES_ESTATICOS.get(nome)
    if versao is None:
        try:
            info = os.stat(os.path.join(app.static_folder, nome))
            versao = format(int(info.st_mtime) ^ info.st_size, "x")
        except OSError:
            versao = ""
        _VERSOES_ESTATICOS[nome] = versao
    if versao:
        valores["v"] = versao


def _validade_estatico(nome_arquivo):
    return 30 * 24 * 3600 if request.args.get("v") else None


app.get_send_file_max_age = _validade_estatico

_TIPOS_COMPRIMIVEIS = {"text/html", "text/css", "text/plain", "application/javascript",
                       "application/json", "application/manifest+json", "image/svg+xml"}


@app.after_request
def _comprimir(resposta):
    """Compacta (gzip) páginas, CSS e dados antes de enviar: a tela de Buscar,
    por exemplo, cai de ~77 KB pra ~17 KB — carrega bem mais rápido no 4G."""
    if (resposta.status_code != 200
            or resposta.mimetype not in _TIPOS_COMPRIMIVEIS
            or "gzip" not in request.headers.get("Accept-Encoding", "").lower()
            or "Content-Encoding" in resposta.headers
            or resposta.is_streamed and not resposta.direct_passthrough):
        return resposta
    resposta.direct_passthrough = False
    dados = resposta.get_data()
    if len(dados) < 1024:
        return resposta
    compactado = gzip.compress(dados, compresslevel=6)
    resposta.set_data(compactado)
    resposta.headers["Content-Encoding"] = "gzip"
    resposta.headers["Content-Length"] = str(len(compactado))
    resposta.vary.add("Accept-Encoding")
    etag, fraco = resposta.get_etag()
    if etag:
        resposta.set_etag(etag + "-gz", weak=fraco)
    return resposta


# Limites por IP (ou por usuário) contra abuso: excesso de cadastros falsos,
# pedidos de senha em massa (que também disparam WhatsApp) e buscas em
# rajada (que poderiam fazer os sites de origem bloquearem o servidor).
_LIMITES = {}
_TRAVA_LIMITES = threading.Lock()


def _passou_do_limite(nome: str, chave: str, maximo: int, janela_seg: int, registrar: bool = True) -> bool:
    """Diz se já passou do limite na janela e (se registrar=True e ainda não
    passou) conta mais uma tentativa."""
    agora = time.time()
    k = (nome, chave)
    with _TRAVA_LIMITES:
        recentes = [t for t in _LIMITES.get(k, []) if agora - t < janela_seg]
        excedeu = len(recentes) >= maximo
        if registrar and not excedeu:
            recentes.append(agora)
        _LIMITES[k] = recentes
        if len(_LIMITES) > 20000:  # faxina pra memória não crescer sem fim
            for velho in [kk for kk, v in _LIMITES.items() if not v or agora - v[-1] > 3600]:
                _LIMITES.pop(velho, None)
    return excedeu


def _ip_visitante() -> str:
    return request.remote_addr or "?"


def _corpo_json():
    """Corpo JSON da requisição como dict (ou {} se vier vazio/inválido)."""
    dados = request.get_json(force=True, silent=True)
    return dados if isinstance(dados, dict) else {}


def _inteiro(valor):
    """Converte pra int ou devolve None (evita erro 500 com id inválido)."""
    try:
        return int(valor)
    except (TypeError, ValueError):
        return None


# Proteção simples contra tentativas de adivinhar senha: depois de várias
# falhas seguidas no mesmo CPF, o login espera alguns minutos.
_TENTATIVAS_LOGIN = {}
_TRAVA_LOGIN = threading.Lock()
MAX_TENTATIVAS_LOGIN = 8
JANELA_LOGIN_SEG = 600


def _chave_login(cpf):
    return "".join(ch for ch in (cpf or "") if ch.isdigit())


def _login_bloqueado(chave):
    agora = time.time()
    with _TRAVA_LOGIN:
        recentes = [t for t in _TENTATIVAS_LOGIN.get(chave, []) if agora - t < JANELA_LOGIN_SEG]
        if recentes:
            _TENTATIVAS_LOGIN[chave] = recentes
        else:
            _TENTATIVAS_LOGIN.pop(chave, None)
        return len(recentes) >= MAX_TENTATIVAS_LOGIN


def _registrar_falha_login(chave):
    agora = time.time()
    with _TRAVA_LOGIN:
        if len(_TENTATIVAS_LOGIN) > 5000:
            for k in [k for k, v in _TENTATIVAS_LOGIN.items() if not v or agora - v[-1] >= JANELA_LOGIN_SEG]:
                _TENTATIVAS_LOGIN.pop(k, None)
        _TENTATIVAS_LOGIN.setdefault(chave, []).append(agora)


def _limpar_falhas_login(chave):
    with _TRAVA_LOGIN:
        _TENTATIVAS_LOGIN.pop(chave, None)


# Aviso pro admin (via CallMeBot) de pedidos de redefinição de senha.
# Dois canais, cada um ligado só se a configuração dele existir:
#   - Telegram: TELEGRAM_USUARIO (ex.: @seunome) — basta ter mandado /start
#     pro @CallMeBot_txtbot no Telegram.
#   - WhatsApp: WHATSAPP_TELEFONE + WHATSAPP_APIKEY.
# Se nenhum estiver configurado, o aviso é só pulado. O envio roda em segundo
# plano: o pedido de senha responde na hora, mesmo se o CallMeBot estiver lento.
def _enviar_avisos_admin(mensagem: str):
    usuario_tg = (os.environ.get("TELEGRAM_USUARIO") or "").strip()
    if usuario_tg:
        if not usuario_tg.startswith("@"):
            usuario_tg = "@" + usuario_tg
        try:
            requests.get(
                "https://api.callmebot.com/text.php",
                params={"user": usuario_tg, "text": mensagem},
                timeout=10,
            )
        except requests.RequestException:
            pass

    telefone = os.environ.get("WHATSAPP_TELEFONE")
    apikey = os.environ.get("WHATSAPP_APIKEY")
    if telefone and apikey:
        try:
            requests.get(
                "https://api.callmebot.com/whatsapp.php",
                params={"phone": telefone, "text": mensagem, "apikey": apikey},
                timeout=10,
            )
        except requests.RequestException:
            pass  # aviso é um "extra" — nunca deve derrubar o pedido de senha


def notificar_admin(mensagem: str):
    if not (os.environ.get("TELEGRAM_USUARIO")
            or (os.environ.get("WHATSAPP_TELEFONE") and os.environ.get("WHATSAPP_APIKEY"))):
        return
    threading.Thread(target=_enviar_avisos_admin, args=(mensagem,), daemon=True).start()


inicializar_auth(app, caminho_dados_persistentes("cannal.db"))

# Senhas novas (cadastro, troca e redefinição) precisam ter pelo menos isso.
# Quem já tem conta com senha menor continua entrando normalmente.
SENHA_MINIMA = 8

MAX_ESPECIALIDADES_POR_BUSCA = 3

ESPECIALIDADES = {
    "neurologista": {"nome": "Neurologista", "doctoralia": "neurologista", "sechat": "n6k3wbvn8k5ajvtb75ikhaaw"},
    "psiquiatra": {"nome": "Psiquiatra", "doctoralia": "psiquiatra", "sechat": "jiagob2yvu3v78j041jb70oz"},
    "neurologista-pediatrico": {"nome": "Neuropediatra", "doctoralia": "neurologista-pediatrico", "sechat": "k9b3og0pgrbrizm138hkamr8"},
    "medico-clinico-geral": {"nome": "Clínico / Clínica Médica", "doctoralia": "medico-clinico-geral", "sechat": "eol7qqtdgtvf8buc2px9hbrq"},
    "pediatra": {"nome": "Pediatra", "doctoralia": "pediatra", "sechat": "dfwzf1l7z9efzucsufeqqeu9"},
    "medico-de-familia": {"nome": "Medicina de Família", "doctoralia": "medico-de-familia", "sechat": "yycwljzp344y327g6rkuyvp1"},
    "especialista-em-dor": {"nome": "Médico da Dor / Anestesiologia", "doctoralia": "especialista-em-dor", "sechat": "zyqcuzojsnjny7wl1lk1g7ah"},
    "ortopedista-traumatologista": {"nome": "Ortopedista", "doctoralia": "ortopedista-traumatologista", "sechat": "wc5z54hdmebdmv8n783876nk"},
    "oncologista": {"nome": "Oncologista", "doctoralia": "oncologista", "sechat": "r742v7goz2k4q6467wsk99di"},
    "geriatra": {"nome": "Geriatra", "doctoralia": "geriatra", "sechat": "ly7dotxntqq446jbqlg6b3l5"},
    "alergista": {"nome": "Alergologista", "doctoralia": "alergista", "sechat": "homqkrplnu5aqeu2uota9wu9"},
    "cardiologista": {"nome": "Cardiologista", "doctoralia": "cardiologista", "sechat": "x14ir5c9kvx9s11bwbvi1mys"},
    "dermatologista": {"nome": "Dermatologista", "doctoralia": "dermatologista", "sechat": "pylctsugyjl97vtfjjs8da1k"},
    "endocrinologista": {"nome": "Endocrinologista", "doctoralia": "endocrinologista", "sechat": "u28uuuaeimpxf3k25cnxru8o"},
    "endocrinologista-pediatrico": {"nome": "Endocrinologista Pediátrico", "doctoralia": "endocrinologista-pediatrico", "sechat": None},
    "especialista-em-medicina-fisica-e-reabilitacao": {"nome": "Medicina Física e Reabilitação", "doctoralia": "especialista-em-medicina-fisica-e-reabilitacao", "sechat": "alfpewwlicu8tux4xcta4yh5"},
    "especialista-em-medicina-preventiva": {"nome": "Medicina Preventiva", "doctoralia": "especialista-em-medicina-preventiva", "sechat": "wutli6r3320ylrt2uclsfp57"},
    "gastroenterologista": {"nome": "Gastroenterologista", "doctoralia": "gastroenterologista", "sechat": "f6w5gv0vbyw6tpambkg1mxag"},
    "generalista": {"nome": "Generalista", "doctoralia": "generalista", "sechat": "c6vsx97abcofv176xjliox2w"},
    "ginecologista": {"nome": "Ginecologista", "doctoralia": "ginecologista", "sechat": "mlhojg8vxi8iftl7jbgudfsw"},
    "medico-acupunturista": {"nome": "Médico Acupunturista", "doctoralia": "medico-acupunturista", "sechat": "djkdf9qhvdw4jzfbhgp4o1pk"},
    "medico-do-esporte": {"nome": "Médico do Esporte", "doctoralia": "medico-do-esporte", "sechat": "kqs3bhoscy5nnf3vdb3vxn1w"},
    "medico-do-sono": {"nome": "Médico do Sono", "doctoralia": "medico-do-sono", "sechat": None},
    "medico-do-trabalho": {"nome": "Médico do Trabalho", "doctoralia": "medico-do-trabalho", "sechat": "evc6v5o06an5xtefc9339ay9"},
    "traumatologista": {"nome": "Traumatologista", "doctoralia": "ortopedista-traumatologista", "sechat": "wc5z54hdmebdmv8n783876nk"},
    "reumatologista": {"nome": "Reumatologista", "doctoralia": "reumatologista", "sechat": "qeiljtcgscrqrwp7wy8jr3cp"},
}

NOMES_ESPECIALIDADES = {slug: info["nome"] for slug, info in ESPECIALIDADES.items()}


@app.route("/cadastro", methods=["GET", "POST"])
def cadastro():
    if usuario_logado():
        return redirect(url_for("index"))

    erro = None
    nome_form = ""
    email_form = ""
    cpf_form = ""
    telefone_form = ""

    if request.method == "POST" and _passou_do_limite("cadastro", _ip_visitante(), 15, 600):
        return render_template("cadastro.html", erro="Muitas tentativas de cadastro seguidas. Aguarde alguns minutos.",
                               nome="", email="", cpf="", telefone=""), 429

    if request.method == "POST":
        nome_form = request.form.get("nome", "").strip()
        email_form = request.form.get("email", "").strip().lower()
        cpf_form = request.form.get("cpf", "").strip()
        telefone_form = request.form.get("telefone", "").strip()
        senha = request.form.get("senha", "")
        confirmar = request.form.get("confirmar", "")

        if not nome_form or not email_form or not cpf_form or not telefone_form or not senha:
            erro = "Preencha todos os campos."
        elif not cpf_valido(cpf_form):
            erro = "CPF inválido. Confira os números digitados."
        elif senha != confirmar:
            erro = "As senhas não coincidem."
        elif len(senha) < SENHA_MINIMA:
            erro = f"A senha precisa ter pelo menos {SENHA_MINIMA} caracteres."
        elif buscar_usuario_por_email(email_form):
            erro = "Já existe uma conta cadastrada com esse e-mail."
        elif buscar_usuario_por_cpf(cpf_form):
            erro = "Já existe uma conta cadastrada com esse CPF."

        if not erro:
            try:
                criar_usuario(nome_form, email_form, cpf_form, senha, telefone=telefone_form)
            except ErroIntegridade:
                erro = "Já existe uma conta cadastrada com esse e-mail ou CPF."
            else:
                # Sem login automático: o cadastro só funciona depois que um
                # admin aprovar (veja a tela de aviso no próprio login.html).
                return redirect(url_for("login", cadastro="sucesso"))

    return render_template("cadastro.html", erro=erro, nome=nome_form, email=email_form, cpf=cpf_form, telefone=telefone_form)


@app.route("/login", methods=["GET", "POST"])
def login():
    if usuario_logado():
        return redirect(url_for("index"))

    erro = None
    if request.method == "POST":
        cpf = request.form.get("cpf", "").strip()
        senha = request.form.get("senha", "")
        chave = _chave_login(cpf)

        if _login_bloqueado(chave) or _passou_do_limite("login_ip", _ip_visitante(), 30, 600, registrar=False):
            erro = "Muitas tentativas de login. Aguarde alguns minutos e tente de novo."
            return render_template("login.html", erro=erro), 429

        usuario = buscar_usuario_por_cpf(cpf)

        if usuario and checar_senha(usuario, senha):
            _limpar_falhas_login(chave)
            if not usuario["aprovado"]:
                erro = "Seu cadastro ainda está sendo analisado pela nossa equipe. Assim que for aprovado, você já poderá entrar."
                return render_template("login.html", erro=erro)
            atualizar_ultimo_login(usuario["id"])
            fazer_login(usuario["id"])
            return redirect(url_for("index"))

        _registrar_falha_login(chave)
        _passou_do_limite("login_ip", _ip_visitante(), 30, 600)
        erro = "CPF ou senha incorretos."

    return render_template("login.html", erro=erro)


def formatar_data_br(valor_iso):
    if not valor_iso:
        return "Nunca"
    try:
        return datetime.fromisoformat(valor_iso).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return valor_iso


@app.route("/perfil", methods=["GET", "POST"])
@login_required
def perfil():
    usuario = usuario_logado()
    erro = None
    sucesso = None

    if request.method == "POST":
        senha_atual = request.form.get("senha_atual", "")
        nova_senha = request.form.get("nova_senha", "")
        confirmar_senha = request.form.get("confirmar_senha", "")

        if not checar_senha(usuario, senha_atual):
            erro = "Senha atual incorreta."
        elif len(nova_senha) < SENHA_MINIMA:
            erro = f"A nova senha precisa ter pelo menos {SENHA_MINIMA} caracteres."
        elif nova_senha != confirmar_senha:
            erro = "As senhas novas não coincidem."
        else:
            definir_senha(usuario["id"], nova_senha)
            usuario = buscar_usuario_por_id(usuario["id"])
            sucesso = "Senha atualizada com sucesso."

    return render_template(
        "perfil.html",
        usuario=usuario,
        cpf_formatado=formatar_cpf(usuario["cpf"]),
        criado_em=formatar_data_br(usuario["criado_em"]),
        ultimo_login=formatar_data_br(usuario["ultimo_login"]),
        erro=erro,
        sucesso=sucesso,
    )


@app.route("/logout")
@login_required
def logout():
    fazer_logout()
    return redirect(url_for("login"))


@app.route("/api/esqueci-senha", methods=["POST"])
def api_esqueci_senha():
    dados = _corpo_json()
    cpf = str(dados.get("cpf", ""))
    # A resposta é sempre "ok" (não revela se o CPF existe). Pedido com CPF
    # inválido, ou repetido demais, só não vira notificação/WhatsApp.
    if not cpf_valido(cpf):
        return jsonify({"ok": True})
    if _passou_do_limite("senha_ip", _ip_visitante(), 5, 600) or \
       _passou_do_limite("senha_cpf", "".join(ch for ch in cpf if ch.isdigit()), 2, 3600):
        return jsonify({"ok": True})
    solicitacao = registrar_solicitacao_senha(cpf)

    nome = solicitacao["nome"] or "CPF não cadastrado no sistema"
    cpf_mostrado = formatar_cpf(solicitacao["cpf"]) if solicitacao["cpf"] else "(não informado)"
    notificar_admin(
        f"🔑 i.cannal — pedido de redefinição de senha\nNome: {nome}\nCPF: {cpf_mostrado}"
    )

    return jsonify({"ok": True})


@app.route("/admin")
@admin_required
def admin():
    usuarios = [
        {
            "id": u["id"],
            "nome": u["nome"],
            "email": u["email"],
            "cpf": formatar_cpf(u["cpf"]),
            "criado_em": formatar_data_br(u["criado_em"]),
            "ultimo_login": formatar_data_br(u["ultimo_login"]),
            "is_admin": bool(u["is_admin"]),
            "eh_mestre": eh_admin_mestre(u),
            "total_painel": u["total_painel"],
        }
        for u in listar_usuarios_com_contagem_painel()
    ]
    return render_template("admin.html", usuarios=usuarios, eh_mestre=eh_admin_mestre(usuario_logado()))


@app.route("/admin/exportar/<int:user_id>")
@admin_required
def admin_exportar_painel(user_id):
    usuario = buscar_usuario_por_id(user_id)
    if not usuario:
        return "Usuário não encontrado.", 404

    painel = listar_painel_usuario(user_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Painel Médico"

    colunas = ["Nome", "CRM", "Especialidade", "Cidade", "UF", "Endereço", "Telefone", "Status", "Adicionado em"]
    ws.append(colunas)

    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill(start_color="343C4C", end_color="343C4C", fill_type="solid")
    for celula in ws[1]:
        celula.font = cabecalho_fonte
        celula.fill = cabecalho_fundo

    for m in painel:
        try:
            data_fmt = datetime.fromisoformat(m["adicionado_em"]).strftime("%d/%m/%Y %H:%M")
        except (ValueError, TypeError):
            data_fmt = m["adicionado_em"] or ""

        ws.append([
            m["nome"], m["crm"], m["especialidade"], m["cidade"], m["uf"],
            m["endereco"], m["telefone"],
            "Visitado" if m["status"] == "visitado" else "Prospectado",
            data_fmt,
        ])

    larguras = [28, 16, 24, 18, 6, 38, 18, 14, 18]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largura

    arquivo = io.BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)

    nome_usuario_arquivo = usuario["nome"].strip().replace(" ", "_")
    nome_arquivo = f"painel_{nome_usuario_arquivo}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return send_file(
        arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nome_arquivo,
    )


@app.route("/admin/painel/<int:user_id>")
@admin_required
def admin_painel_usuario(user_id):
    usuario = buscar_usuario_por_id(user_id)
    if not usuario:
        return jsonify({"erro": "Usuário não encontrado."}), 404

    painel = listar_painel_usuario(user_id)
    prospectados, agendados, visitados, revisitar = bucket_painel(painel)

    agendados_dict = [dict(m) for m in agendados]
    for m in agendados_dict:
        m["compromisso"] = buscar_compromisso_do_medico(user_id, m["id"])

    return jsonify({
        "usuario": {"nome": usuario["nome"], "email": usuario["email"]},
        "prospectados": [dict(m) for m in prospectados],
        "agendados": agendados_dict,
        "visitados": [dict(m) for m in visitados],
        "revisitar": [dict(m) for m in revisitar],
    })


@app.route("/admin/estatisticas/<int:user_id>")
@admin_required
def admin_estatisticas_usuario(user_id):
    usuario = buscar_usuario_por_id(user_id)
    if not usuario:
        return jsonify({"erro": "Usuário não encontrado."}), 404

    stats = estatisticas_usuario(user_id)
    stats["usuario"] = {"nome": usuario["nome"], "email": usuario["email"]}
    stats["medicos_removidos"] = [dict(m) for m in listar_medicos_removidos_usuario(user_id)]
    return jsonify(stats)


@app.route("/admin/notificacoes")
@admin_required
def admin_notificacoes():
    pendentes = listar_solicitacoes_senha_pendentes()
    cadastros = listar_cadastros_pendentes()
    return jsonify({
        "total": len(pendentes) + len(cadastros),
        "solicitacoes": [
            {
                "id": s["id"],
                "cpf": formatar_cpf(s["cpf"]),
                "nome": s["nome"],
                "email": s["email"],
                "criado_em": formatar_data_br(s["criado_em"]),
            }
            for s in pendentes
        ],
        "cadastros_pendentes": [
            {
                "id": c["id"],
                "nome": c["nome"],
                "email": c["email"],
                "cpf": formatar_cpf(c["cpf"]),
                "telefone": c["telefone"],
                "criado_em": formatar_data_br(c["criado_em"]),
            }
            for c in cadastros
        ],
    })


@app.route("/admin/cadastros-pendentes/<int:usuario_id>/aprovar", methods=["POST"])
@admin_required
def admin_cadastro_aprovar(usuario_id):
    if not buscar_usuario_por_id(usuario_id):
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404
    ok = aprovar_usuario(usuario_id)
    return jsonify({"ok": ok})


@app.route("/admin/notificacoes/<int:solicitacao_id>/atender", methods=["POST"])
@admin_required
def admin_notificacao_atender(solicitacao_id):
    ok = marcar_solicitacao_atendida(solicitacao_id)
    return jsonify({"ok": ok})


@app.route("/admin/notificacoes/<int:solicitacao_id>/redefinir", methods=["POST"])
@admin_required
def admin_notificacao_redefinir(solicitacao_id):
    dados = _corpo_json()
    nova_senha = dados.get("senha", "")

    if len(nova_senha) < SENHA_MINIMA:
        return jsonify({"ok": False, "mensagem": f"A senha precisa ter pelo menos {SENHA_MINIMA} caracteres."}), 400

    solicitacao = buscar_solicitacao_senha_por_id(solicitacao_id)
    if not solicitacao:
        return jsonify({"ok": False, "mensagem": "Solicitação não encontrada."}), 404

    usuario = buscar_usuario_por_cpf(solicitacao["cpf"])
    if not usuario:
        return jsonify({"ok": False, "mensagem": "Não existe conta cadastrada com esse CPF."}), 404

    definir_senha(usuario["id"], nova_senha)
    marcar_solicitacao_atendida(solicitacao_id)
    return jsonify({"ok": True})


@app.route("/admin/dashboard")
@admin_required
def admin_dashboard():
    limpar_medicos_inativos_todos_usuarios()  # mantém em dia mesmo quem não loga há um tempo
    return render_template(
        "dashboard.html",
        stats=estatisticas_dashboard(),
        medicos_removidos_geral=listar_medicos_removidos_geral(),
    )


@app.route("/")
@login_required
def index():
    usuario = usuario_logado()
    limpar_medicos_inativos(usuario["id"])  # roda a cada visita — é barato e mantém o painel em dia
    dias_inativo = dias_sem_visitar(usuario["id"])
    visitas_amanha = contar_visitas_amanha(usuario["id"])
    return render_template(
        "index.html",
        especialidades=NOMES_ESPECIALIDADES,
        max_especialidades=MAX_ESPECIALIDADES_POR_BUSCA,
        aviso_inatividade=(dias_inativo is not None and dias_inativo >= 10),
        dias_inativo=dias_inativo,
        visitas_amanha=visitas_amanha,
        resultados_visita=RESULTADOS_VISITA,
        notificacoes_nao_lidas=contar_notificacoes_nao_lidas(usuario["id"]),
    )


@app.route("/api/buscar")
@login_required
def api_buscar():
    especialidades_selecionadas = request.args.getlist("especialidade")
    cidade = request.args.get("cidade", "")[:80]
    uf = request.args.get("uf", "")[:2]

    if _passou_do_limite("busca", str(usuario_logado()["id"]), 40, 300):
        return jsonify({"erro": "Muitas buscas seguidas. Aguarde um minuto e tente de novo."}), 429

    if not especialidades_selecionadas:
        return jsonify({"erro": "Selecione ao menos uma especialidade."}), 400
    if len(especialidades_selecionadas) > MAX_ESPECIALIDADES_POR_BUSCA:
        return jsonify({"erro": f"Selecione no máximo {MAX_ESPECIALIDADES_POR_BUSCA} especialidades por busca."}), 400
    for slug in especialidades_selecionadas:
        if slug not in ESPECIALIDADES:
            return jsonify({"erro": f"Especialidade inválida: {slug}"}), 400
    if not cidade.strip():
        return jsonify({"erro": "Informe a cidade."}), 400

    usuario = usuario_logado()
    registrar_busca(usuario["id"])

    nomes_especialidades = [ESPECIALIDADES[slug]["nome"] for slug in especialidades_selecionadas]
    rotulo_especialidades = ", ".join(nomes_especialidades)

    # Monta a lista de tarefas a rodar em paralelo:
    # - Doctoralia e Sechat entram uma vez POR especialidade selecionada
    #   (cada um tem uma URL diferente por especialidade).
    # - As fontes genéricas (Ama-me, Kaya Doc, Cannaceia) não filtram por
    #   especialidade — a página é a mesma não importa o que você busca —
    #   então elas entram só UMA VEZ na lista toda.
    tarefas = []  # cada item: (rótulo_pra_erro, função, args)

    for slug in especialidades_selecionadas:
        info = ESPECIALIDADES[slug]
        nome_especialidade = info["nome"]

        tarefas.append((
            f"Doctoralia ({nome_especialidade})",
            buscar_medicos_doctoralia,
            (info["doctoralia"], nome_especialidade, cidade, uf),
        ))

        if info["sechat"]:
            tarefas.append((
                f"Sechat ({nome_especialidade})",
                buscar_medicos_sechat,
                (info["sechat"], nome_especialidade, cidade, uf),
            ))

        tarefas.append((
            f"APEPI ({nome_especialidade})",
            buscar_medicos_apepi,
            (slug, nome_especialidade, cidade, uf),
        ))

    for chave_fonte, dados_fonte in FONTES_GENERICAS.items():
        tarefas.append((
            dados_fonte["nome"],
            buscar_medicos_fonte_generica,
            (chave_fonte, rotulo_especialidades, cidade, uf),
        ))

    # A base própria (médicos importados de planilha) é só uma consulta no
    # banco — sempre ao vivo, pra uma importação nova aparecer na hora.
    todos_medicos = buscar_base_propria(especialidades_selecionadas, cidade, uf, NOMES_ESPECIALIDADES)

    # Mesma busca feita há pouco (por qualquer pessoa da equipe): reaproveita
    # o que os sites devolveram em vez de consultar tudo de novo.
    chave_cache = (tuple(sorted(especialidades_selecionadas)), normalizar_texto_busca(cidade), uf.strip().upper())
    achou, guardado = _CACHE_BUSCAS.pegar(chave_cache)
    if achou:
        resultados_sites, fontes_com_erro = [dict(m) for m in guardado], []
    else:
        resultados_sites, fontes_com_erro = [], []
        # Roda tudo em paralelo (até 12 buscas ao mesmo tempo) em vez de uma
        # atrás da outra — é isso que evita o erro de timeout quando várias
        # especialidades são selecionadas juntas.
        with ThreadPoolExecutor(max_workers=12) as executor:
            futuros = {
                executor.submit(funcao, *args): rotulo
                for rotulo, funcao, args in tarefas
            }
            for futuro in as_completed(futuros):
                rotulo = futuros[futuro]
                try:
                    resultados_sites.extend(futuro.result())
                except Exception as e:
                    app.logger.warning("Fonte com erro na busca — %s: %s", rotulo, e)
                    fontes_com_erro.append(f"{rotulo}: fonte indisponível no momento")
        if not fontes_com_erro:  # só guarda resultado completo
            _CACHE_BUSCAS.guardar(chave_cache, [dict(m) for m in resultados_sites])

    todos_medicos = _juntar_duplicados(todos_medicos + resultados_sites)

    # Médico que já está no painel do usuário logado: mostra "Presente no
    # painel médico" em vez do botão de adicionar. Médico que já está no
    # painel de OUTRO usuário: nem aparece — já tem alguém de olho nele.
    chaves_meu_painel = chaves_no_painel_usuario(usuario["id"])
    chaves_outros_paineis = chaves_no_painel_todos() - chaves_meu_painel

    medicos_visiveis = []
    for medico in todos_medicos:
        chave = computar_chave_medico(medico)
        if chave in chaves_outros_paineis:
            continue
        medico["no_painel"] = chave in chaves_meu_painel
        medicos_visiveis.append(medico)
    todos_medicos = medicos_visiveis

    return jsonify({
        "especialidades": nomes_especialidades,
        "cidade": cidade,
        "uf": uf,
        "total": len(todos_medicos),
        "medicos": todos_medicos,
        "avisos": fontes_com_erro,
    })


_CACHE_BUSCAS = CacheTTL(10 * 60, 300)  # buscas iguais em até 10 min


def normalizar_texto_busca(texto: str) -> str:
    return " ".join(normalizar(texto or "").split())


def _juntar_duplicados(medicos: list) -> list:
    """O mesmo médico (mesmo CRM) pode vir de mais de uma fonte — mostra uma
    vez só, completando telefone/endereço que faltarem com o que a outra
    fonte trouxe. Sem CRM, não junta (nome igual pode ser outra pessoa)."""
    vazios = {"", "Não encontrado", "Não disponível", None}
    por_chave = {}
    saida = []
    for m in medicos:
        chave = computar_chave_medico(m)
        if not chave.startswith("crm:"):
            saida.append(m)
            continue
        if chave not in por_chave:
            por_chave[chave] = m
            saida.append(m)
            continue
        principal = por_chave[chave]
        for campo in ("telefone", "endereco", "perfil_url"):
            if principal.get(campo) in vazios and m.get(campo) not in vazios:
                principal[campo] = m[campo]
    return saida


@app.route("/api/painel", methods=["GET"])
@login_required
def api_painel_listar():
    usuario = usuario_logado()
    painel = listar_painel_usuario(usuario["id"])
    prospectados, agendados, visitados, revisitar = bucket_painel(painel)

    agendados_dict = [dict(m) for m in agendados]
    for m in agendados_dict:
        m["compromisso"] = buscar_compromisso_do_medico(usuario["id"], m["id"])

    return jsonify({
        "prospectados": [dict(m) for m in prospectados],
        "agendados": agendados_dict,
        "visitados": [dict(m) for m in visitados],
        "revisitar": [dict(m) for m in revisitar],
    })


@app.route("/api/painel/adicionar", methods=["POST"])
@login_required
def api_painel_adicionar():
    usuario = usuario_logado()
    medico = _corpo_json()

    if not medico.get("nome"):
        return jsonify({"ok": False, "mensagem": "Dados do médico incompletos."}), 400

    ok, mensagem = adicionar_ao_painel(usuario["id"], medico)
    return jsonify({"ok": ok, "mensagem": mensagem})


@app.route("/api/painel/<int:entrada_id>/remover", methods=["POST"])
@login_required
def api_painel_remover(entrada_id):
    usuario = usuario_logado()
    ok = remover_medico_painel(usuario["id"], entrada_id)
    return jsonify({"ok": ok})


@app.route("/api/notificacoes")
@login_required
def api_notificacoes():
    usuario = usuario_logado()
    notificacoes = listar_notificacoes_usuario(usuario["id"])
    return jsonify({"notificacoes": [dict(n) for n in notificacoes]})


@app.route("/api/notificacoes/marcar-lidas", methods=["POST"])
@login_required
def api_notificacoes_marcar_lidas():
    usuario = usuario_logado()
    marcar_notificacoes_lidas(usuario["id"])
    return jsonify({"ok": True})


@app.route("/api/painel/mover", methods=["POST"])
@login_required
def api_painel_mover():
    usuario = usuario_logado()
    dados = _corpo_json()
    entrada_id = _inteiro(dados.get("id"))
    novo_status = dados.get("status")
    resultado = dados.get("resultado")

    # "agendado" só existe através do botão Agendar (que exige data e horário).
    if not entrada_id or novo_status not in ("prospectado", "visitado"):
        return jsonify({"ok": False, "mensagem": "Requisição inválida."}), 400
    if resultado is not None:
        if novo_status != "visitado":
            return jsonify({"ok": False, "mensagem": "Resultado só se aplica ao marcar como visitado."}), 400
        if resultado not in RESULTADOS_VISITA:
            return jsonify({"ok": False, "mensagem": "Resultado da visita inválido."}), 400

    ok = mover_no_painel(usuario["id"], entrada_id, novo_status, resultado=resultado)
    return jsonify({"ok": ok})


@app.route("/api/painel/<int:entrada_id>/observacoes", methods=["POST"])
@login_required
def api_painel_observacoes(entrada_id):
    usuario = usuario_logado()
    dados = _corpo_json()
    texto = (dados.get("texto") or "").strip()

    ok = salvar_observacoes(usuario["id"], entrada_id, texto)
    return jsonify({"ok": ok})


@app.route("/api/painel/agendar", methods=["POST"])
@login_required
def api_painel_agendar():
    usuario = usuario_logado()
    dados = _corpo_json()
    entrada_id = _inteiro(dados.get("id"))
    data = str(dados.get("data") or "").strip()
    horario = str(dados.get("horario") or "").strip()

    if not entrada_id or not data or not horario:
        return jsonify({"ok": False, "mensagem": "Escolha uma data e um horário."}), 400

    ok, mensagem = agendar_medico(usuario["id"], entrada_id, data, horario)
    return jsonify({"ok": ok, "mensagem": mensagem})


@app.route("/api/painel/<int:entrada_id>/cancelar-agendamento", methods=["POST"])
@login_required
def api_painel_cancelar_agendamento(entrada_id):
    usuario = usuario_logado()
    ok = cancelar_agendamento(usuario["id"], entrada_id)
    return jsonify({"ok": ok})


@app.route("/agenda")
@login_required
def agenda():
    return render_template("agenda.html")


@app.route("/api/agenda/mes")
@login_required
def api_agenda_mes():
    usuario = usuario_logado()
    ano = _inteiro(request.args.get("ano"))
    mes = _inteiro(request.args.get("mes"))
    if ano is None or mes is None or not (2000 <= ano <= 2100) or not (1 <= mes <= 12):
        return jsonify({"erro": "Parâmetros inválidos."}), 400

    return jsonify({"dias": listar_compromissos_mes(usuario["id"], ano, mes)})


@app.route("/api/agenda/dia")
@login_required
def api_agenda_dia():
    usuario = usuario_logado()
    data = (request.args.get("data") or "").strip()
    if not data_valida(data):
        return jsonify({"erro": "Informe uma data válida (AAAA-MM-DD)."}), 400

    return jsonify({"horarios": listar_horarios_dia(usuario["id"], data)})


@app.route("/api/roteiro")
@login_required
def api_roteiro():
    usuario = usuario_logado()
    data = (request.args.get("data") or "").strip()
    if not data_valida(data):
        return jsonify({"erro": "Informe uma data válida (AAAA-MM-DD)."}), 400

    return jsonify({"visitas": roteiro_do_dia(usuario["id"], data)})


@app.route("/admin/base-propria")
@admin_required
def admin_base_propria():
    especialidades_ordenadas = sorted(NOMES_ESPECIALIDADES.items(), key=lambda item: item[1])
    return render_template(
        "base_propria.html",
        total_base=contar_base_propria(),
        especialidades=especialidades_ordenadas,
    )


@app.route("/admin/base-propria/importar", methods=["POST"])
@admin_required
def admin_base_propria_importar():
    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify({"ok": False, "mensagem": "Selecione um arquivo CSV."}), 400
    if not arquivo.filename.lower().endswith(".csv"):
        return jsonify({"ok": False, "mensagem": "O arquivo precisa ser um .csv."}), 400

    bruto = arquivo.read(10 * 1024 * 1024)  # limite de 10 MB
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            texto = bruto.decode("latin-1")
        except UnicodeDecodeError:
            return jsonify({"ok": False, "mensagem": "Não foi possível ler esse arquivo. Salve como CSV (UTF-8) e tente de novo."}), 400

    try:
        amostra = texto[:2000]
        separador = ";" if amostra.count(";") > amostra.count(",") else ","
        leitor = csv.DictReader(io.StringIO(texto), delimiter=separador)
        leitor.fieldnames = [(nome or "").strip().lower() for nome in (leitor.fieldnames or [])]
        colunas_obrigatorias = {"nome", "especialidade"}
        if not colunas_obrigatorias.issubset(set(leitor.fieldnames)):
            return jsonify({
                "ok": False,
                "mensagem": "O CSV precisa ter as colunas 'nome' e 'especialidade'. Colunas aceitas: nome, crm, especialidade, cidade, uf, endereco, telefone.",
            }), 400
        linhas = list(leitor)
    except csv.Error:
        return jsonify({"ok": False, "mensagem": "Não foi possível ler esse CSV. Confira o formato do arquivo."}), 400

    if not linhas:
        return jsonify({"ok": False, "mensagem": "O arquivo está vazio."}), 400

    usuario = usuario_logado()
    resultado = importar_base_propria(usuario["id"], linhas, set(ESPECIALIDADES.keys()))
    return jsonify({"ok": True, **resultado})


@app.route("/admin/usuarios/<int:user_id>/painel/adicionar", methods=["POST"])
@admin_required
def admin_painel_adicionar(user_id):
    if not buscar_usuario_por_id(user_id):
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404

    medico = _corpo_json()
    if not medico.get("nome"):
        return jsonify({"ok": False, "mensagem": "Dados do médico incompletos."}), 400

    ok, mensagem = adicionar_ao_painel(user_id, medico)
    return jsonify({"ok": ok, "mensagem": mensagem})


@app.route("/admin/usuarios/<int:user_id>/painel/<int:entrada_id>/remover", methods=["POST"])
@admin_required
def admin_painel_remover(user_id, entrada_id):
    if not buscar_usuario_por_id(user_id):
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404

    ok = remover_medico_painel(user_id, entrada_id)
    return jsonify({"ok": ok})


@app.route("/admin/usuarios/<int:user_id>/painel/importar", methods=["POST"])
@admin_required
def admin_painel_importar(user_id):
    if not buscar_usuario_por_id(user_id):
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404

    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify({"ok": False, "mensagem": "Selecione um arquivo CSV."}), 400
    if not arquivo.filename.lower().endswith(".csv"):
        return jsonify({"ok": False, "mensagem": "O arquivo precisa ser um .csv."}), 400

    bruto = arquivo.read(10 * 1024 * 1024)  # limite de 10 MB
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            texto = bruto.decode("latin-1")
        except UnicodeDecodeError:
            return jsonify({"ok": False, "mensagem": "Não foi possível ler esse arquivo. Salve como CSV (UTF-8) e tente de novo."}), 400

    try:
        amostra = texto[:2000]
        separador = ";" if amostra.count(";") > amostra.count(",") else ","
        leitor = csv.DictReader(io.StringIO(texto), delimiter=separador)
        leitor.fieldnames = [(nome or "").strip().lower() for nome in (leitor.fieldnames or [])]
        if "nome" not in leitor.fieldnames:
            return jsonify({
                "ok": False,
                "mensagem": "O CSV precisa ter uma coluna 'nome'. Colunas aceitas: nome, crm, especialidade, cidade, uf, endereco, telefone.",
            }), 400
        linhas = list(leitor)
    except csv.Error:
        return jsonify({"ok": False, "mensagem": "Não foi possível ler esse CSV. Confira o formato do arquivo."}), 400

    if not linhas:
        return jsonify({"ok": False, "mensagem": "O arquivo está vazio."}), 400

    resultado = admin_importar_painel_usuario(user_id, linhas)
    return jsonify({"ok": True, **resultado})


@app.route("/admin/usuarios/<int:user_id>/tornar-admin", methods=["POST"])
@admin_required
def admin_tornar_admin(user_id):
    if not buscar_usuario_por_id(user_id):
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404

    ok = tornar_admin(user_id)
    return jsonify({"ok": ok})


@app.route("/admin/usuarios/<int:user_id>/remover-admin", methods=["POST"])
@admin_required
def admin_remover_admin(user_id):
    usuario_atual = usuario_logado()
    if not eh_admin_mestre(usuario_atual):
        return jsonify({"ok": False, "mensagem": "Só o admin principal pode remover o acesso de outros admins."}), 403

    alvo = buscar_usuario_por_id(user_id)
    if not alvo:
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404
    if alvo["id"] == usuario_atual["id"]:
        return jsonify({"ok": False, "mensagem": "Você não pode remover o próprio acesso de admin."}), 400

    ok = remover_admin(user_id)
    return jsonify({"ok": ok})


@app.route("/admin/usuarios/<int:user_id>/excluir", methods=["POST"])
@admin_required
def admin_excluir_usuario(user_id):
    usuario_atual = usuario_logado()
    alvo = buscar_usuario_por_id(user_id)
    if not alvo:
        return jsonify({"ok": False, "mensagem": "Usuário não encontrado."}), 404
    if alvo["id"] == usuario_atual["id"]:
        return jsonify({"ok": False, "mensagem": "Você não pode excluir a sua própria conta."}), 400
    if eh_admin_mestre(alvo):
        return jsonify({"ok": False, "mensagem": "A conta do admin principal não pode ser excluída."}), 403
    if alvo["is_admin"] and not eh_admin_mestre(usuario_atual):
        return jsonify({"ok": False, "mensagem": "Só o admin principal pode excluir a conta de outro admin."}), 403

    ok = excluir_usuario(user_id)
    return jsonify({"ok": ok})


@app.route("/sw.js")
def service_worker():
    # Precisa vir da raiz (não de /static/), senão o navegador só deixa ele
    # controlar as páginas dentro de /static/ — e o objetivo é o site inteiro.
    resposta = send_from_directory(resource_path("static"), "sw.js")
    resposta.headers["Content-Type"] = "application/javascript"
    resposta.headers["Service-Worker-Allowed"] = "/"
    resposta.headers["Cache-Control"] = "no-cache"
    return resposta


@app.route("/offline")
def offline():
    return render_template("offline.html")


def abrir_navegador():
    webbrowser.open("http://127.0.0.1:5000")


if __name__ == "__main__":
    threading.Timer(1.2, abrir_navegador).start()
    app.run(debug=False, use_reloader=False, port=5000)
